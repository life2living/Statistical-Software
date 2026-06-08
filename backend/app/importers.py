from __future__ import annotations

import csv
from datetime import date, datetime
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook


def parse_tabular_file(filename: str, content: bytes) -> list[dict[str, Any]]:
    lower_name = filename.lower()
    if lower_name.endswith(".csv"):
        return parse_csv(content)
    if lower_name.endswith(".xlsx") or lower_name.endswith(".xlsm"):
        return parse_xlsx(content)
    raise ValueError("Only CSV, XLSX, and XLSM files are supported.")


def parse_csv(content: bytes) -> list[dict[str, Any]]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(StringIO(text))
    return [coerce_row(row) for row in reader]


def parse_xlsx(content: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        return []

    names = [str(header).strip() if header is not None else f"Column {index + 1}" for index, header in enumerate(headers)]
    parsed_rows: list[dict[str, Any]] = []
    for raw_row in rows:
        row = {name: raw_row[index] if index < len(raw_row) else None for index, name in enumerate(names)}
        if any(value is not None for value in row.values()):
            parsed_rows.append(coerce_row(row))
    return parsed_rows


def coerce_row(row: dict[str, Any]) -> dict[str, Any]:
    return {str(key).strip(): coerce_value(value) for key, value in row.items()}


def coerce_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, int | float):
        return value
    if not isinstance(value, str):
        return str(value)

    stripped = value.strip()
    if stripped == "":
        return None
    try:
        if any(marker in stripped for marker in [".", "e", "E"]):
            return float(stripped)
        return int(stripped)
    except ValueError:
        return stripped
